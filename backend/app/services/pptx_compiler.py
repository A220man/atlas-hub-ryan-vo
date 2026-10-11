import io
from typing import List, Optional
from pptx import Presentation as PptxPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

from app.services.slide_synthesizer import SlideGraphNode


class PPTXCompiler:
    THEME_COLORS = {
        "modern-navy": {"primary": RGBColor(14, 30, 64), "secondary": RGBColor(0, 114, 206), "accent": RGBColor(230, 92, 0), "text": RGBColor(33, 37, 41)},
        "emerald-clean": {"primary": RGBColor(10, 61, 46), "secondary": RGBColor(16, 149, 107), "accent": RGBColor(245, 158, 11), "text": RGBColor(20, 24, 22)},
        "slate-minimal": {"primary": RGBColor(30, 41, 59), "secondary": RGBColor(71, 85, 105), "accent": RGBColor(99, 102, 241), "text": RGBColor(15, 23, 42)},
    }

    def ingest_template(self, file_bytes: bytes, filename: str) -> dict:
        prs = PptxPresentation(io.BytesIO(file_bytes))
        ratio = "16:9" if (prs.slide_width / max(1, prs.slide_height)) >= 1.5 else "4:3"
        return {"filename": filename, "layout_count": len(prs.slide_layouts), "aspect_ratio": ratio}

    def compile(
        self,
        slide_nodes: List[SlideGraphNode],
        template_id: str = "modern-navy",
        aspect_ratio: str = "16:9",
        custom_template_bytes: Optional[bytes] = None,
    ) -> bytes:
        if custom_template_bytes:
            prs = PptxPresentation(io.BytesIO(custom_template_bytes))
            for r_id in [prs.slides._sldIdLst[i].rId for i in range(len(prs.slides))]:
                prs.part.drop_rel(r_id)
            prs.slides._sldIdLst.clear()
            layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
            colors = self.THEME_COLORS["modern-navy"]
        else:
            prs = PptxPresentation()
            prs.slide_width = Inches(13.333 if aspect_ratio == "16:9" else 10.0)
            prs.slide_height = Inches(7.5)
            layout = prs.slide_layouts[6]
            colors = self.THEME_COLORS.get(template_id, self.THEME_COLORS["modern-navy"])

        for node in slide_nodes:
            slide = prs.slides.add_slide(layout)
            if node.speaker_notes:
                slide.notes_slide.notes_text_frame.text = node.speaker_notes

            shapes = [self._render_header(slide, node, colors)]
            if node.layout_type == "metrics_chart":
                shapes.extend(self._render_metrics_chart(slide, node, colors))
            elif node.layout_type == "comparison_grid":
                shapes.extend(self._render_comparison_grid(slide, node, colors))
            else:
                shapes.extend(self._render_standard_bullets(slide, node, colors))

            self._apply_slide_transition(slide, node.content.transition or "fade")
            sp_ids = [s.shape_id for s in shapes if hasattr(s, "shape_id")]
            if sp_ids:
                self._apply_shape_animations(slide, sp_ids)

        out_stream = io.BytesIO()
        prs.save(out_stream)
        return out_stream.getvalue()

    def _apply_slide_transition(self, slide, trans_type: str) -> None:
        t_xml = {
            "push": f'<p:transition {nsdecls("p")} spd="med"><p:push dir="l"/></p:transition>',
            "wipe": f'<p:transition {nsdecls("p")} spd="med"><p:wipe dir="r"/></p:transition>',
            "split": f'<p:transition {nsdecls("p")} spd="med"><p:split orient="horz"/></p:transition>',
        }.get(trans_type, f'<p:transition {nsdecls("p")} spd="med"><p:fade/></p:transition>')
        slide._element.append(parse_xml(t_xml))

    def _apply_shape_animations(self, slide, shape_ids: List[int]) -> None:
        pars = "".join(
            f'<p:par><p:cTn id="{10+i}" fill="hold"><p:stCondLst><p:cond delay="{i*200}"/></p:stCondLst>'
            f'<p:childTn><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{100+i}" dur="400"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTn></p:cTn></p:par>'
            for i, spid in enumerate(shape_ids[:5])
        )
        timing_xml = (
            f'<p:timing {nsdecls("p")}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="always" nodeType="tmRoot">'
            f'<p:childTn><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            f'<p:childTn>{pars}</p:childTn></p:cTn>'
            f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            f'</p:seq></p:childTn></p:cTn></p:par></p:tnLst></p:timing>'
        )
        slide._element.append(parse_xml(timing_xml))

    def _render_header(self, slide, node: SlideGraphNode, colors: dict):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = node.title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = colors["primary"]
        if node.content.header and node.content.header != node.title:
            p2 = tf.add_paragraph()
            p2.text = node.content.header.upper()
            p2.font.size = Pt(11)
            p2.font.bold = True
            p2.font.color.rgb = colors["secondary"]
        return box

    def _render_standard_bullets(self, slide, node: SlideGraphNode, colors: dict) -> list:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.5))
        tf = box.text_frame
        tf.word_wrap = True
        for idx, bullet in enumerate(node.content.bullets):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            p.text = f"•  {bullet}"
            p.font.size = Pt(16)
            p.font.color.rgb = colors["text"]
            p.space_after = Pt(12)
        shapes = [box]
        if node.content.callout:
            cbox = slide.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.9))
            cp = cbox.text_frame.paragraphs[0]
            cp.text = f"Takeaway: {node.content.callout}"
            cp.font.size = Pt(13)
            cp.font.bold = True
            cp.font.color.rgb = colors["accent"]
            shapes.append(cbox)
        return shapes

    def _render_metrics_chart(self, slide, node: SlideGraphNode, colors: dict) -> list:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.2), Inches(4.5))
        tf = box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = node.content.metric_value or "94.6%"
        p1.font.size = Pt(38)
        p1.font.bold = True
        p1.font.color.rgb = colors["accent"]
        p1.space_after = Pt(8)
        for b in node.content.bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.size = Pt(13)
            pb.font.color.rgb = colors["text"]
        series_data = node.content.chart_series or [{"category": "Baseline", "value": 50.0}, {"category": "Pipeline", "value": 94.6}]
        cdata = CategoryChartData()
        cdata.categories = [item["category"] for item in series_data]
        cdata.add_series("Metric Score", [float(item["value"]) for item in series_data])
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.4), Inches(1.8), Inches(7.1), Inches(4.5), cdata)
        return [box, chart]

    def _render_comparison_grid(self, slide, node: SlideGraphNode, colors: dict) -> list:
        cols = node.content.table_columns or ["Aspect", "Baseline", "Atlas Hub"]
        rows = node.content.table_rows or [["Integrity Gate", "Silent Exit 1", "Actionable Diagnostics"]]
        t_shape = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(0.8), Inches(1.8), Inches(11.7), Inches(3.8))
        table = t_shape.table
        for c_idx, c_name in enumerate(cols):
            cell = table.cell(0, c_idx)
            cell.text = c_name
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = colors["primary"]
        for r_idx, r_vals in enumerate(rows):
            for c_idx, val in enumerate(r_vals):
                if c_idx < len(cols):
                    cell = table.cell(r_idx + 1, c_idx)
                    cell.text = str(val)
                    for p in cell.text_frame.paragraphs:
                        p.font.size = Pt(12)
                        p.font.color.rgb = colors["text"]
        return [t_shape]
