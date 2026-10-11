import io
from pptx import Presentation as PptxPresentation
from app.services.document_parser import DocumentParser
from app.services.slide_synthesizer import SlideSynthesizer
from app.services.pptx_compiler import PPTXCompiler


def test_pptx_compiler_generates_valid_presentation():
    """Verify that pptx compiler generates valid OpenXML binary with shapes and notes."""
    parser = DocumentParser()
    synthesizer = SlideSynthesizer()
    compiler = PPTXCompiler()

    doc_text = (
        "# Native Deck Synthesis\n\n"
        "## Performance Metrics\n"
        "- Throughput is 88 tokens/s.\n"
        "- Memory footprint is 140 mb.\n\n"
        "## Feature Comparison\n"
        "- Span verification ensures factual grounding.\n"
    )

    doc_result = parser.parse(doc_text)
    slides = synthesizer.synthesize(doc_result)

    pptx_bytes = compiler.compile(slides, template_id="modern-navy", aspect_ratio="16:9")

    # Assert binary output begins with standard ZIP/OpenXML magic number
    assert len(pptx_bytes) > 1000
    assert pptx_bytes.startswith(b"PK\x03\x04")

    # Load back using python-pptx to verify integrity
    loaded_prs = PptxPresentation(io.BytesIO(pptx_bytes))
    assert len(loaded_prs.slides) == len(slides)

    # Check that speaker notes and shapes are embedded
    for idx, slide in enumerate(loaded_prs.slides):
        assert len(slide.shapes) > 0
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text
            assert len(notes_text) > 0


def test_pptx_compiler_supports_alternative_templates():
    """Verify compilation with emerald-clean and 4:3 aspect ratio."""
    parser = DocumentParser()
    synthesizer = SlideSynthesizer()
    compiler = PPTXCompiler()

    doc_result = parser.parse("# Topic Brief\n- Key takeaway item.")
    slides = synthesizer.synthesize(doc_result)

    pptx_bytes = compiler.compile(slides, template_id="emerald-clean", aspect_ratio="4:3")
    assert len(pptx_bytes) > 1000
    loaded = PptxPresentation(io.BytesIO(pptx_bytes))
    assert len(loaded.slides) >= 3
